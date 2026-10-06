from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader


@dataclass(frozen=True)
class FormatExtractionResult:
    extracted_text: str
    extracted_data: dict[str, Any]
    metadata: dict[str, Any]
    provider: str
    warnings: list[str]


class DocumentFormatError(Exception):
    def __init__(self, message: str, *, code: str = "INVALID_FILE") -> None:
        super().__init__(message)
        self.code = code


def _business_fields(text: str) -> dict[str, str]:
    fields: dict[str, str] = {}
    patterns = {
        "remittance_number": r"(?:remittance(?:\s+number)?|remittance_number)\s*[:=]\s*([A-Z0-9-]+)",
        "customer_id": r"customer(?:\s+id)?\s*[:=]\s*([A-Z0-9-]+)",
        "invoice_number": r"invoice(?:\s+number)?\s*[:=]\s*([A-Z0-9-]+)",
    }
    for name, pattern in patterns.items():
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            fields[name] = match.group(1)
    return fields


extract_business_fields = _business_fields


class PDFDocumentParser:
    provider = "pypdf"

    def parse(self, path: Path) -> FormatExtractionResult:
        try:
            reader = PdfReader(str(path))
            pages: list[str] = []
            warnings: list[str] = []
            for page in reader.pages:
                pages.append(page.extract_text() or "")
            text = "\n".join(pages).strip()
        except Exception as error:
            raise DocumentFormatError(f"unable to parse PDF: {error}") from error
        if not text:
            warnings.append("text extraction was insufficient; OCR would be required")
        return FormatExtractionResult(text, _business_fields(text), {"page_count": len(reader.pages), "text_sufficient": bool(text)}, self.provider, warnings)


class XLSXDocumentParser:
    provider = "openpyxl"

    def parse(self, path: Path) -> FormatExtractionResult:
        try:
            workbook = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=False, keep_links=False)
        except Exception as error:
            raise DocumentFormatError(f"unable to parse workbook: {error}") from error
        lines: list[str] = []
        sources: list[dict[str, Any]] = []
        extracted_data: dict[str, Any] = {}
        sheet_names = list(workbook.sheetnames)
        for worksheet in workbook.worksheets:
            header_values: list[str] | None = None
            for row in worksheet.iter_rows(values_only=False):
                values = [cell.value for cell in row]
                if not any(value is not None for value in values):
                    continue
                if header_values is None:
                    header_values = [str(value).strip().lower() for value in values if value is not None]
                elif len(header_values) == len(values):
                    for header, value in zip(header_values, values):
                        if value is not None:
                            if "remittance" in header:
                                extracted_data["remittance_number"] = str(value)
                            elif "customer" in header:
                                extracted_data["customer_id"] = str(value)
                rendered = " | ".join(str(value) for value in values if value is not None)
                lines.append(f"[{worksheet.title}] {rendered}")
                sources.append({"sheet": worksheet.title, "cells": [cell.coordinate for cell in row if cell.value is not None]})
        workbook.close()
        text = "\n".join(lines)
        return FormatExtractionResult(text, {**_business_fields(text), **extracted_data}, {"sheet_names": sheet_names, "sources": sources[:100], "text_sufficient": bool(text)}, self.provider, [])


class DocumentParser:
    def parse(self, path: Path, *, mime_type: str, document_name: str) -> FormatExtractionResult:
        suffix = Path(document_name).suffix.lower()
        if mime_type == "application/pdf" or suffix == ".pdf":
            return PDFDocumentParser().parse(path)
        if suffix in {".xlsx", ".xlsm"} or mime_type in {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/vnd.ms-excel.sheet.macroEnabled.12"}:
            return XLSXDocumentParser().parse(path)
        raise DocumentFormatError(f"unsupported document format: {document_name}", code="UNSUPPORTED_FILE")
